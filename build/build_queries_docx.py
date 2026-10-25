"""
Builds build/gdocs/Northwind_Queries.docx: the four SQL files with a plain-language
explanation before each one, every file starting on a new page, and each step's
code in its own shaded box. Upload to Drive and open with Google Docs.
"""
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "build" / "gdocs" / "Northwind_Queries.docx"

RULES = [
    ("R1", "sessions", 'Same channel written 12 ways ("direct ", "Display ")', "Map to 7 channel names", "1,284"),
    ("R2", "sessions", "135 URL spellings of ~20 pages", "Clean the URL, group into 7 page types", "all"),
    ("R3", "sessions", "Visit with no visitor id", "Keep it, mark is_visitor_missing", "300"),
    ("R4", "signups", "Exact duplicate rows", "Remove", "66"),
    ("R5", "signups", "Fake company sizes (−5, 999,999)", "Set to unknown", "30"),
    ("R6", "signups", "Same lead twice with different details", "Keep the most complete row", "4 leads"),
    ("R7", "signups", "Signup dated July, after the data ends", "Use the date of its visit", "8"),
    ("R8", "signups", "Points to a visit that doesn't exist", "Keep it, mark is_session_orphan", "25"),
    ("R9", "opportunities", "Exact duplicate rows", "Remove", "18"),
    ("R10", "opportunities", "Amount 1,000× the price (35,000 / 110,000)", "Divide by 1,000", "6"),
    ("R11", "opportunities", "$0 deals (nobody pays)", "Keep, mark is_zero_amount; never a win", "30"),
    ("R12", "opportunities", "Negative days_to_close (sign bug)", "Recompute from the dates", "10"),
    ("R13", "opportunities", "Points to a lead that doesn't exist", "Keep it, mark is_lead_orphan", "15"),
    ("R14", "signups", "Country written as the text 'N/A'", "Set to unknown", "25"),
]

FILES = [
    ("00_load_raw.sql", "Load the data (local only)", [
        ("What it does", "turns the three CSV files into tables called sessions, signups and opportunities."),
        ("Why", "in Snowflake these tables already exist, so you skip this file. On a laptop with DuckDB, this step creates them with the same names, so every other query runs unchanged in both places."),
    ]),
    ("01_clean.sql", "Clean the tables and connect them", [
        ("What it does", "builds four clean tables, one after the other. Each one uses the one before."),
        ("Idea 1: LEFT JOIN", "keeps every row on the left even when there's no match on the right. Broken rows get marked instead of silently disappearing."),
        ("Idea 2: flags are 1/0", "so SUM counts them and AVG gives a rate. is_win is the most important one: 1 = paying customer within 30 days, 0 = not."),
    ]),
    ("02_funnel.sql", "The funnel numbers", [
        ("What it does", "turns signup_outcomes into the tables behind every number in the analysis."),
        ("Conversion rate", "AVG(is_win). The average of a column of 1s and 0s is the share of 1s, which is the conversion rate."),
    ]),
    ("03_decomposition.sql", "Why the number of paying customers changed", [
        ("The question", "paying customers went from 316 to 168 per month. How much of that is because of more signups, a different SKU mix, or worse conversion inside each SKU?"),
        ("The idea", "wins = signups × conversion, and conversion = each SKU's share × its own conversion. Change one thing at a time, holding the rest at the average of both periods, and the three pieces add up exactly to the total change."),
        ("The answer", "volume +56, mix −201, rate −3 paying customers per month. Almost all of the drop is the SKU mix."),
        ("The one assumption", "Free CRM didn't exist in Q1, so it has no Q1 conversion. We use its Q2 conversion for both periods, so its whole effect counts as mix."),
    ]),
]

STEP_NOTES = {
    "sessions_clean": "Fixes channel names (R1), turns 135 messy URLs into 7 page types (R2), marks visits with no visitor (R3).",
    "signups_clean": "Removes duplicates (R4, R6), blanks out fake values (R5, R14), fixes July dates using the visit date (R7), marks signups whose visit doesn't exist (R8).",
    "opportunities_clean": "Removes duplicates (R9), fixes amounts 1,000× too high (R10), marks $0 deals (R11), recomputes days to close (R12), marks deals whose lead doesn't exist (R13).",
    "signup_outcomes": "The table everything else reads: where the signup came from, which SKU, and is_win.",
}

BANNER = re.compile(r"^-- -{20,}\n((?:-- (?!-).*\n)+)-- -{20,}\n", re.M)


def split_steps(sql_text):
    body = sql_text.split("-- " + "=" * 69 + "\n")[-1].strip() + "\n"
    out, pos, title, note = [], 0, None, None
    for m in BANNER.finditer(body):
        chunk = body[pos:m.start()].strip()
        if chunk:
            out.append((title, note, chunk))
        lines = [l[3:].strip() for l in m.group(1).strip("\n").split("\n")]
        title, note = lines[0], " ".join(lines[1:]) or None
        pos = m.end()
    chunk = body[pos:].strip()
    if chunk:
        out.append((title, note, chunk))
    return out


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def code_box(doc, code):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    shade(cell, "F3F3F1")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    lines = code.split("\n")
    for i, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "Courier New"; r.font.size = Pt(8.5)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        if line.lstrip().startswith("--"):
            r.font.color.rgb = RGBColor(0x5F, 0x6B, 0x5F)
        if i < len(lines) - 1:
            r.add_break()
    doc.add_paragraph()


def bullets(doc, items):
    for label, text in items:
        p = doc.add_paragraph(style="List Bullet")
        b = p.add_run(label + ": "); b.bold = True
        p.add_run(text)


doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2)

doc.add_heading("Northwind CRM — Queries", 0)
doc.add_paragraph("Snowflake SQL. The same files also run unchanged in DuckDB on a laptop. Run them in order: "
                  "00 (local only), 01, 02, 03. Source tables in Snowflake: sessions, signups, opportunities. "
                  "Repo: github.com/northwind-acv/plg-conversion-panel")
doc.add_heading("What we had to fix so the data is usable", 1)
doc.add_paragraph("All fixes happen in 01_clean.sql; 02 and 03 only read the clean tables. "
                  "Rule ids match the comments in the code and the Data methodology doc.")
tbl = doc.add_table(rows=1, cols=5); tbl.style = "Table Grid"
for i, h in enumerate(["Rule", "Table", "Problem found", "What we did", "Rows"]):
    c = tbl.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; shade(c, "E8EEF6")
for row in RULES:
    cells = tbl.add_row().cells
    for i, v in enumerate(row):
        cells[i].text = v
for row in tbl.rows:
    for c in row.cells:
        for p in c.paragraphs:
            for r in p.runs: r.font.size = Pt(9)

for fname, title, notes in FILES:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(f"{fname} — {title}", 1)
    bullets(doc, notes)
    steps = split_steps((ROOT / "sql" / fname).read_text())
    for step_title, note, code in steps:
        if step_title:
            doc.add_heading(step_title, 3)
            key = next((k for k in STEP_NOTES if k in step_title), None)
            text = " ".join(x for x in [note, STEP_NOTES.get(key)] if x)
            if text:
                doc.add_paragraph(text)
        code_box(doc, code)

doc.save(OUT)
print("wrote", OUT)
